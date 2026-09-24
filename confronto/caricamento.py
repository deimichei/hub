"""
Caricamento dei file del cliente nelle tabelle di staging (schema "giacenza").

Fa il lavoro che nella versione psql faceva \\copy, ma:
  * legge gli Excel (.xlsx) direttamente: niente conversione a mano in CSV
  * abbina le colonne per NOME, non per posizione: se Zoho sposta o
    aggiunge colonne l'import non si rompe; se ne manca una che serve,
    si ferma con un messaggio chiaro
  * trova da solo la riga d'intestazione (il Batch di Zoho ha un titolo sopra)

I nomi delle colonne vengono normalizzati con la STESSA regola usata per
generare lo schema (00_schema.sql): minuscolo, senza accenti, spazi e simboli
-> "_", parole riservate con "_" finale, nomi ripetuti con _2, _3.
"""

import csv
import io
import re
import unicodedata
import warnings
from datetime import date, datetime

from openpyxl import load_workbook

# Gli export di Zoho non hanno lo stile predefinito: openpyxl lo segnala
# con un avviso innocuo, che qui nascondiamo
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

# Parole riservate SQL: come in 00_schema.sql ricevono un "_" finale
RISERVATE = {
    'all', 'and', 'any', 'array', 'as', 'asc', 'both', 'case', 'cast', 'check',
    'collate', 'column', 'constraint', 'create', 'current_date', 'default',
    'desc', 'distinct', 'do', 'else', 'end', 'except', 'false', 'for',
    'foreign', 'from', 'grant', 'group', 'having', 'in', 'initially',
    'intersect', 'into', 'leading', 'limit', 'not', 'null', 'offset', 'on',
    'only', 'or', 'order', 'placing', 'primary', 'references', 'returning',
    'select', 'some', 'symmetric', 'table', 'then', 'to', 'trailing', 'true',
    'union', 'unique', 'user', 'using', 'when', 'where', 'window', 'with',
}

# I 4 documenti Excel: tabella di staging e colonne che DEVONO esserci
# (quelle usate dagli script di pulizia). Servono anche a riconoscere
# la riga d'intestazione.
DOCUMENTI = {
    'batch': {
        'tabella': 'import_batch_full',
        'obbligatorie': ['sku', 'batch_number', 'balance_quantity'],
    },
    'sales_order': {
        'tabella': 'import_sales_order_full',
        'obbligatorie': ['salesorder_number', 'line_item_location_name'],
    },
    'package': {
        'tabella': 'import_package_full',
        'obbligatorie': ['so_number', 'status', 'sku', 'batch_reference', 'quantity_out'],
    },
    'giacenza': {
        'tabella': 'import_giacenza_full',
        'obbligatorie': ['descrizione', 'numero_lotto', 'quantita', 'bloccato'],
    },
}

# I 3 file di riferimento (CSV mantenuti a mano, separatore virgola)
RIFERIMENTI = {
    'sku_alias': {'tabella': 'ref_sku_alias', 'obbligatorie': ['zoho_sku', 'sku_canonico']},
    'sku_old': {'tabella': 'ref_sku_old', 'obbligatorie': ['no_sku', 'motivo']},
    'batch_notes': {'tabella': 'ref_batch_notes', 'obbligatorie': ['sku', 'batch', 'note']},
}


class ErroreFile(Exception):
    """Errore con messaggio pensato per l'utente (mostrato nella pagina)."""


def normalizza(nomi):
    """Intestazioni del file -> nomi di colonna come in 00_schema.sql."""
    risultato, visti = [], {}
    for nome in nomi:
        n = unicodedata.normalize('NFKD', str(nome or '')).encode('ascii', 'ignore').decode()
        n = re.sub(r'[^0-9a-zA-Z]+', '_', n.strip()).strip('_').lower() or 'col'
        if n in RISERVATE:
            n += '_'
        if n in visti:
            visti[n] += 1
            n = f'{n}_{visti[n]}'
        else:
            visti[n] = 1
        risultato.append(n)
    return risultato


def come_testo(valore):
    """Cella Excel -> testo per la staging (tutte le colonne sono text)."""
    if valore is None:
        return None
    if isinstance(valore, bool):
        return str(valore)
    if isinstance(valore, float) and valore.is_integer():
        return str(int(valore))          # 672.0 -> "672"
    if isinstance(valore, (datetime, date)):
        return valore.isoformat()[:10]   # data -> "2026-09-24"
    testo = str(valore)
    return testo if testo != '' else None


def colonne_tabella(cursore, tabella):
    cursore.execute(
        "SELECT column_name FROM information_schema.columns "
        "WHERE table_schema = 'giacenza' AND table_name = %s "
        "ORDER BY ordinal_position",
        [tabella],
    )
    return [r[0] for r in cursore.fetchall()]


def _trova_intestazione(righe, obbligatorie, nome_file):
    """Cerca nelle prime 20 righe quella che contiene le colonne obbligatorie."""
    for i, riga in enumerate(righe[:20]):
        nomi = normalizza(riga)
        if all(c in nomi for c in obbligatorie):
            return i, nomi
    raise ErroreFile(
        f"{nome_file}: non trovo la riga d'intestazione con le colonne "
        f"{', '.join(obbligatorie)}. E' il file giusto?"
    )


def _scrivi(cursore, tabella, nomi, righe, nome_file):
    """Scrive le righe nella tabella con COPY (veloce), solo le colonne note."""
    colonne_db = colonne_tabella(cursore, tabella)
    indici = [(i, n) for i, n in enumerate(nomi) if n in colonne_db]
    colonne = ', '.join(n for _, n in indici)
    cursore.execute(f'TRUNCATE {tabella}')
    n_righe = 0
    # cursore.cursor = il cursore psycopg sotto quello di Django
    with cursore.cursor.copy(f'COPY {tabella} ({colonne}) FROM STDIN') as copy:
        for riga in righe:
            valori = [come_testo(riga[i]) if i < len(riga) else None for i, _ in indici]
            if all(v is None or str(v).strip() == '' for v in valori):
                continue                  # riga vuota (es. in fondo al file)
            copy.write_row(valori)
            n_righe += 1
    if n_righe == 0:
        raise ErroreFile(f'{nome_file}: il file non contiene righe di dati.')
    return n_righe


def carica_excel(cursore, chiave, file_caricato):
    """Uno dei 4 documenti (.xlsx) -> tabella import_*_full."""
    doc = DOCUMENTI[chiave]
    try:
        wb = load_workbook(file_caricato, read_only=True, data_only=True)
    except Exception:
        raise ErroreFile(f'{file_caricato.name}: non riesco ad aprirlo come file Excel (.xlsx).')
    righe = [list(r) for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    i, nomi = _trova_intestazione(righe, doc['obbligatorie'], file_caricato.name)
    return _scrivi(cursore, doc['tabella'], nomi, righe[i + 1:], file_caricato.name)


def carica_csv(cursore, chiave, file_caricato):
    """Uno dei 3 file di riferimento (.csv, separatore virgola) -> tabella ref_*."""
    rif = RIFERIMENTI[chiave]
    testo = file_caricato.read().decode('utf-8-sig')
    righe = list(csv.reader(io.StringIO(testo), delimiter=','))
    if not righe:
        raise ErroreFile(f'{file_caricato.name}: il file e\' vuoto.')
    i, nomi = _trova_intestazione(righe, rif['obbligatorie'], file_caricato.name)
    # le chiavi vengono tolte prima: la normalizzazione (01_riferimenti.sql)
    # puo' creare doppioni temporanei, che lo script poi elimina
    cursore.execute(f"ALTER TABLE {rif['tabella']} DROP CONSTRAINT IF EXISTS {rif['tabella']}_pkey")
    return _scrivi(cursore, rif['tabella'], nomi, righe[i + 1:], file_caricato.name)
